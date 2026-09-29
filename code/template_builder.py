"""
template_builder.py — Study list -> studyname_master workbook bytes.

The only application module (besides test_meta_pipeline.py) that imports
openpyxl. Everything is built in memory (io.BytesIO); there are no
filesystem writes (CLAUDE.md rule 3). No streamlit import.

PHASE2_ARCHITECTURE.md section 1.9 and section 6.

FINDING (reported, not silently worked around — see the Stage 2 handback):
section 6 item 2 says a study name's XML-illegal characters are "removed
from the written value" by openpyxl itself. Verified empirically against
openpyxl 3.1.5: assigning such a string to a cell raises
`IllegalCharacterError` instead of silently stripping it. `_strip_illegal_xml_chars`
below does the stripping ourselves, before the value ever reaches openpyxl,
so template generation cannot crash on a pathological master Study_Name
(edge case P2-29 still fires downstream in Stage 4, unaffected: the merge
key differs before and after stripping either way).
"""

from __future__ import annotations

import io
from typing import Sequence

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font, Protection
from openpyxl.worksheet.worksheet import Worksheet

import config
from schema import build_column_index, normalize_column


def list_template_studies(master_df) -> list[str]:
    """P5: distinct Study_Name values, VERBATIM (exact-string distinct),
    first-appearance order, values whose .strip() == "" excluded.

    Study_Name is found via build_column_index — the same normalized-name
    lookup every other Phase 2 module uses (P24).
    """
    column_index = build_column_index(list(master_df.columns))
    study_column = column_index[normalize_column(config.STUDY_NAME_COL)]

    seen: set[str] = set()
    studies: list[str] = []
    for value in master_df[study_column]:
        if value.strip() == "":
            continue
        if value not in seen:
            seen.add(value)
            studies.append(value)
    return studies


def build_template_bytes(study_names: Sequence[str]) -> bytes:
    """Section 6. Workbook built entirely in memory; wb.save(io.BytesIO())."""
    workbook = Workbook()
    _build_template_sheet(workbook, study_names)
    _build_glossary_sheet(workbook)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _strip_illegal_xml_chars(value: str) -> str:
    """See the module docstring FINDING. ILLEGAL_CHARACTERS_RE is openpyxl's
    own pattern; this only pre-empts the exception it would otherwise raise.
    """
    return ILLEGAL_CHARACTERS_RE.sub("", value)


def _set_column_range(
    worksheet: Worksheet,
    key: str,
    min_col: int,
    max_col: int,
    width: int,
    locked: bool,
    percent: bool = False,
) -> None:
    dimension = worksheet.column_dimensions[key]
    dimension.min = min_col
    dimension.max = max_col
    dimension.width = width
    dimension.protection = Protection(locked=locked)
    if percent:
        dimension.number_format = config.PCT_NUMBER_FORMAT


def _build_template_sheet(workbook: Workbook, study_names: Sequence[str]) -> None:
    worksheet = workbook.active
    worksheet.title = config.TEMPLATE_SHEET_TITLE
    headers = config.TEMPLATE_HEADERS
    last_column = len(headers)          # column G
    percent_column = 4                  # column D — Pct_HH_Buying

    # Row 1: bold header, column A locked (default), B..G unlocked.
    for col_idx, header in enumerate(headers, start=1):
        cell = worksheet.cell(row=1, column=col_idx, value=header)
        cell.font = Font(bold=True)
        if col_idx > 1:
            cell.protection = Protection(locked=False)

    # Rows 2..n+1: column A = study names (forced to string data), B..G
    # materialized empty with Protection(locked=False); D also gets the
    # percent number format.
    for row_offset, name in enumerate(study_names):
        row_idx = row_offset + 2
        name_cell = worksheet.cell(row=row_idx, column=1, value=_strip_illegal_xml_chars(name))
        name_cell.data_type = "s"
        for col_idx in range(2, last_column + 1):
            value_cell = worksheet.cell(row=row_idx, column=col_idx, value=None)
            value_cell.protection = Protection(locked=False)
            if col_idx == percent_column:
                value_cell.number_format = config.PCT_NUMBER_FORMAT

    worksheet.freeze_panes = "A2"

    # Column dimensions: cells typed later anywhere in B onward, including
    # H+ headers and values, stay unlocked (section 6 item 4).
    worksheet.column_dimensions["A"].width = config.TEMPLATE_STUDY_COL_WIDTH
    _set_column_range(worksheet, "B", 2, 3, config.TEMPLATE_VALUE_COL_WIDTH, locked=False)
    _set_column_range(
        worksheet, "D", 4, 4, config.TEMPLATE_VALUE_COL_WIDTH, locked=False, percent=True
    )
    _set_column_range(
        worksheet, "E", 5, config.XLSX_MAX_COLUMN, config.TEMPLATE_VALUE_COL_WIDTH, locked=False
    )

    # Sheet protection: on, NO password, formatColumns allowed (design
    # default 18) so users may widen columns. Every other flag stays at
    # openpyxl's default.
    worksheet.protection.sheet = True
    worksheet.protection.formatColumns = False  # DESIGN DEFAULT — pending confirmation, see PHASE2_ARCHITECTURE.md section 11, item 18


def _build_glossary_sheet(workbook: Workbook) -> None:
    """Sheet 2, not protected. Content verbatim from config (P31)."""
    worksheet = workbook.create_sheet(config.GLOSSARY_SHEET_TITLE)

    title_cell = worksheet.cell(row=1, column=1, value=config.GLOSSARY_FIELDS_TITLE)
    title_cell.font = Font(bold=True)

    for col_idx, header in enumerate(config.GLOSSARY_FIELD_COLUMNS, start=1):
        cell = worksheet.cell(row=2, column=col_idx, value=header)
        cell.font = Font(bold=True)

    what_to_enter_column = 3   # "What to enter" — the long, wrapped column (C)
    field_rows_start = 3
    for row_offset, field_row in enumerate(config.GLOSSARY_FIELD_ROWS):
        row_idx = field_rows_start + row_offset
        for col_idx, cell_value in enumerate(field_row, start=1):
            cell = worksheet.cell(row=row_idx, column=col_idx, value=cell_value)
            cell.data_type = "s"

    blank_after_fields = field_rows_start + len(config.GLOSSARY_FIELD_ROWS)   # row 10
    extras_title_row = blank_after_fields + 1                                # row 11
    extras_title_cell = worksheet.cell(row=extras_title_row, column=1, value=config.GLOSSARY_EXTRAS_TITLE)
    extras_title_cell.font = Font(bold=True)
    worksheet.cell(row=extras_title_row + 1, column=1, value=config.GLOSSARY_EXTRAS_TEXT)   # row 12

    rules_title_row = extras_title_row + 3   # row 14 (row 13 left blank)
    rules_title_cell = worksheet.cell(row=rules_title_row, column=1, value=config.GLOSSARY_RULES_TITLE)
    rules_title_cell.font = Font(bold=True)
    for rule_offset, rule in enumerate(config.GLOSSARY_RULES):
        worksheet.cell(row=rules_title_row + 1 + rule_offset, column=1, value=f"- {rule}")

    worksheet.column_dimensions["A"].width = 18
    worksheet.column_dimensions["B"].width = 28
    worksheet.column_dimensions["C"].width = 80
    worksheet.column_dimensions["D"].width = 20
    for row_offset in range(len(config.GLOSSARY_FIELD_ROWS)):
        cell = worksheet.cell(row=field_rows_start + row_offset, column=what_to_enter_column)
        cell.alignment = cell.alignment.copy(wrap_text=True)
